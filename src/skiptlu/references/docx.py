from __future__ import annotations

from pathlib import Path
import zipfile

from docx import Document

from skiptlu.references.models import (
    DocumentImage,
    DocumentParagraph,
    DocumentSection,
    DocumentRequirementHints,
    DocumentTable,
    ParsedDocument,
)


class DocxReferenceParser:
    """Extracts structural facts from a Moodle-provided DOCX without interpreting its meaning."""

    def parse(self, path: Path) -> ParsedDocument:
        doc = Document(str(path))
        paragraphs: list[DocumentParagraph] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else None
            level = self._heading_level(style, text)
            paragraphs.append(
                DocumentParagraph(text=text, style=style, level=level)
            )

        captions = [p.text.strip() for p in doc.paragraphs if self._is_caption(p.text, p.style.name if p.style else None)]
        sections = [self._section_info(section) for section in doc.sections]
        has_footnotes = self._has_part(path, "word/footnotes.xml")
        has_toc = self._has_toc(path)

        tables = [
            DocumentTable(
                rows=[
                    [cell.text.strip() for cell in row.cells]
                    for row in table.rows
                ]
            )
            for table in doc.tables
        ]

        images = self._images(path, doc)

        return ParsedDocument(
            path=path,
            file_type="docx",
            paragraphs=paragraphs,
            tables=tables,
            images=images,
            sections=sections,
            hints=DocumentRequirementHints(
                headings=[p.text for p in paragraphs if p.level is not None],
                captions=captions,
                has_toc=has_toc,
                has_footnotes=has_footnotes,
                has_landscape_section=any(s.orientation == "landscape" for s in sections),
            ),
            metadata={
                "section_count": len(doc.sections),
                "inline_shape_count": len(doc.inline_shapes),
            },
        )

    @staticmethod
    def _is_caption(text: str, style: str | None) -> bool:
        lowered = text.strip().lower()
        return (style or "").lower() == "caption" or lowered.startswith(("ảnh ", "bảng ", "hình ", "table ", "figure "))

    @staticmethod
    def _section_info(section) -> DocumentSection:
        width = section.page_width.twips if section.page_width else None
        height = section.page_height.twips if section.page_height else None
        orientation = "landscape" if width and height and width > height else "portrait"
        margins = {
            "top": section.top_margin.twips,
            "bottom": section.bottom_margin.twips,
            "left": section.left_margin.twips,
            "right": section.right_margin.twips,
        }
        return DocumentSection(
            orientation=orientation,
            page_width_twips=width,
            page_height_twips=height,
            margins_twips=margins,
        )

    @staticmethod
    def _has_part(path: Path, name: str) -> bool:
        with zipfile.ZipFile(path) as archive:
            return name in archive.namelist()

    @staticmethod
    def _has_toc(path: Path) -> bool:
        with zipfile.ZipFile(path) as archive:
            if "word/document.xml" not in archive.namelist():
                return False
            xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
            return "TOC" in xml or "w:fldSimple" in xml and "TOC" in xml

    @staticmethod
    def _heading_level(style: str | None, text: str) -> int | None:
        if style:
            lowered = style.lower()
            if lowered.startswith("heading 1"):
                return 1
            if lowered.startswith("heading 2"):
                return 2
            if lowered.startswith("heading 3"):
                return 3

        normalized = text.lower()
        if normalized.startswith("chương "):
            return 1
        return None

    @staticmethod
    def _images(path: Path, doc: Document) -> list[DocumentImage]:
        images: list[DocumentImage] = []
        with zipfile.ZipFile(path) as archive:
            media = sorted(
                name for name in archive.namelist()
                if name.startswith("word/media/")
            )

        for index, name in enumerate(media, start=1):
            shape = doc.inline_shapes[index - 1] if index <= len(doc.inline_shapes) else None
            images.append(
                DocumentImage(
                    name=Path(name).name,
                    width=shape.width if shape else None,
                    height=shape.height if shape else None,
                )
            )
        return images
