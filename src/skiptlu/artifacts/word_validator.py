from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.section import WD_ORIENT

from skiptlu.models.word_task import WordTaskSpecification


def _clean_heading(value: str) -> str:
    import re
    return re.sub(
        r"^CHƯƠNG\s+\d+\.\s*|^\d+(?:\.\d+){1,3}\.\s*",
        "",
        value,
    ).strip()


def _caption_text(caption: str) -> str:
    return caption.split(".", 2)[-1].strip() if "." in caption else caption


class WordValidator:
    """Validate a generated DOCX against its task specification."""

    def __init__(self, opener: Callable[[str], Document] = Document):
        self._opener = opener

    def validate(self, path: Path, spec: WordTaskSpecification) -> list[str]:
        try:
            doc = self._opener(str(path))
        except Exception as exc:
            return [f"Invalid DOCX: {exc}"]

        errors: list[str] = []
        errors.extend(self._validate_content(doc, spec))
        errors.extend(self._validate_structure(doc, spec))
        errors.extend(self._validate_package(path, spec))
        return errors

    def _validate_content(self, doc: Document, spec: WordTaskSpecification) -> list[str]:
        errors: list[str] = []
        text = "\n".join(p.text for p in doc.paragraphs)

        for level in (1, 2, 3):
            for expected in spec.derived_headings(level):
                normalized = _clean_heading(expected)
                if normalized and normalized not in text:
                    errors.append(f"Missing heading content: {normalized}")

        for caption in spec.derived_captions():
            normalized = _caption_text(caption)
            if normalized and normalized not in text:
                errors.append(f"Missing caption text: {normalized}")

        return errors

    def _validate_structure(self, doc: Document, spec: WordTaskSpecification) -> list[str]:
        errors: list[str] = []
        content = spec.content_specification
        template_mode = bool(
            spec.task_specification is not None
            and str(spec.task_specification.metadata.get("artifact_strategy", "")) == "use_template"
        )

        # In template mode, structure is validated against the template
        # snapshot instead of imposing the legacy generated-document shape.
        if not template_mode and len(doc.sections) < 3:
            errors.append("Missing section breaks for Chapter 1 and landscape section")

        if content and content.include_landscape_section and not template_mode:
            if not any(s.orientation == WD_ORIENT.LANDSCAPE for s in doc.sections):
                errors.append("Missing landscape section required by specification")

        if content and content.tables and len(doc.tables) < len(content.tables):
            errors.append(
                f"Expected at least {len(content.tables)} tables, found {len(doc.tables)}"
            )

        if content and content.figures and len(doc.inline_shapes) < len(content.figures):
            errors.append(
                f"Expected at least {len(content.figures)} figures, found {len(doc.inline_shapes)}"
            )

        return errors

    def _validate_package(self, path: Path, spec: WordTaskSpecification) -> list[str]:
        errors: list[str] = []
        content = spec.content_specification
        needs_footnotes = bool(content and content.include_footnotes)

        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())

                if needs_footnotes and "word/footnotes.xml" not in names:
                    errors.append("Missing footnotes.xml")

                xml = archive.read("word/document.xml").decode("utf-8")

                if content and content.include_toc and "TOC" not in xml:
                    errors.append("Missing TOC field")

                if needs_footnotes and "footnoteReference" not in xml:
                    errors.append("Missing footnote reference")
        except (OSError, zipfile.BadZipFile, KeyError) as exc:
            errors.append(f"Invalid DOCX package: {exc}")

        return errors


def validate_word_artifact(path: Path, spec: WordTaskSpecification) -> list[str]:
    """Compatibility wrapper for the legacy validator API."""
    return WordValidator().validate(path, spec)
