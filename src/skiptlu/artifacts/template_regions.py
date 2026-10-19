from __future__ import annotations

"""Deterministic structural fingerprinting and semantic region mapping."""

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
import re

from docx import Document


@dataclass(frozen=True)
class TemplateRegion:
    id: str
    kind: str
    index: int
    anchor: str
    paragraph_indices: tuple[int, ...] = ()
    table_index: int | None = None


@dataclass(frozen=True)
class TemplateFingerprint:
    source: str
    digest: str
    paragraph_count: int
    table_count: int
    section_count: int
    regions: tuple[TemplateRegion, ...]


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().casefold()


def _digest(doc: Document) -> str:
    parts: list[str] = []
    for i, paragraph in enumerate(doc.paragraphs):
        parts.append(f"P:{i}:{paragraph.style.name}:{_normalize(paragraph.text)}")
    for i, table in enumerate(doc.tables):
        parts.append(f"T:{i}:{len(table.rows)}x{len(table.columns)}")
        for row in table.rows:
            parts.append("|".join(_normalize(cell.text) for cell in row.cells))
    for i, section in enumerate(doc.sections):
        parts.append(
            f"S:{i}:{section.orientation}:{section.page_width}:{section.page_height}:"
            f"{section.top_margin}:{section.bottom_margin}:{section.left_margin}:{section.right_margin}"
        )
    return sha256("\n".join(parts).encode("utf-8")).hexdigest()


def fingerprint_document(doc: Document, source: str = "<memory>") -> TemplateFingerprint:
    regions: list[TemplateRegion] = []

    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        normalized = _normalize(text)
        if not normalized:
            continue

        if normalized.startswith("sinh viên thực hiện"):
            regions.append(TemplateRegion(
                "cover.student_name", "student_info", index, text, (index,)
            ))
        elif normalized.startswith("mã sinh viên"):
            regions.append(TemplateRegion(
                "cover.student_id", "student_info", index, text, (index,)
            ))
        elif normalized.startswith("ngành:"):
            regions.append(TemplateRegion(
                "cover.major", "student_info", index, text, (index,)
            ))
        elif normalized in {"mục lục", "danh mục minh họa"}:
            regions.append(TemplateRegion(
                f"front.{normalized.replace(' ', '_')}", "front_matter", index, text, (index,)
            ))
        elif paragraph.style.name.startswith("Heading 1"):
            n = sum(r.kind == "heading1" for r in regions)
            regions.append(TemplateRegion(f"body.heading1.{n}", "heading1", index, text, (index,)))
        elif paragraph.style.name.startswith("Heading 2"):
            n = sum(r.kind == "heading2" for r in regions)
            regions.append(TemplateRegion(f"body.heading2.{n}", "heading2", index, text, (index,)))
        elif paragraph.style.name.startswith("Heading 3"):
            n = sum(r.kind == "heading3" for r in regions)
            regions.append(TemplateRegion(f"body.heading3.{n}", "heading3", index, text, (index,)))

    for table_index, table in enumerate(doc.tables):
        anchor = " ".join(
            cell.text.strip()
            for row in table.rows[:1]
            for cell in row.cells
            if cell.text.strip()
        )
        regions.append(
            TemplateRegion(
                f"table.{table_index}", "table", table_index, anchor, (), table_index
            )
        )

    return TemplateFingerprint(
        source=source,
        digest=_digest(doc),
        paragraph_count=len(doc.paragraphs),
        table_count=len(doc.tables),
        section_count=len(doc.sections),
        regions=tuple(regions),
    )


def fingerprint_docx(path: Path) -> TemplateFingerprint:
    return fingerprint_document(Document(str(path)), str(path))


def fingerprint_payload(fingerprint: TemplateFingerprint) -> dict[str, object]:
    return {
        "source": fingerprint.source,
        "digest": fingerprint.digest,
        "paragraph_count": fingerprint.paragraph_count,
        "table_count": fingerprint.table_count,
        "section_count": fingerprint.section_count,
        "regions": [
            {
                "id": region.id,
                "kind": region.kind,
                "index": region.index,
                "anchor": region.anchor,
                "paragraph_indices": list(region.paragraph_indices),
                "table_index": region.table_index,
            }
            for region in fingerprint.regions
        ],
    }
