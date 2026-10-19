from __future__ import annotations

from pathlib import Path
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

from skiptlu.models.artifact import Artifact, ArtifactKind
from skiptlu.models.word_task import WordTaskSpecification
from skiptlu.tasks.layout import LayoutSpecification, default_word_layout
from skiptlu.artifacts.word_backend import NativeWordBackend, WordBackend
from skiptlu.artifacts.word_validator import WordValidator
from skiptlu.artifacts.word_support import (
    add_field as _field,
    caption as _caption,
    clear_paragraph as _clear,
    font as _font,
    footnote_reference as _footnote_reference,
    inject_footnote as _inject_footnote,
    insert_section_before as _insert_section_before,
    page_number as _page_number,
    paragraph_after as _after,
    replace_paragraph as _replace,
    update_fields as _update_fields,
)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _legacy_font_removed(run, size=13, bold=False, italic=False):
    return _font(run, size, bold, italic)


def _clear(p):
    return __import__("skiptlu.artifacts.word_support", fromlist=["clear_paragraph"]).clear_paragraph(p)


def _replace(p, text):
    return __import__("skiptlu.artifacts.word_support", fromlist=["replace_paragraph"]).replace_paragraph(p, text)


def _spacing(p, before=6, after=6, line=1.3):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def _field(p, instruction, placeholder="1"):
    r = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, separate, t, end])
    _font(r)


def _after(p, style=None):
    new_p = OxmlElement("w:p")
    p._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    new = Paragraph(new_p, p._parent)
    if style:
        new.style = style
    return new


def _style(doc, name, base="Normal"):
    if name in doc.styles:
        return doc.styles[name]
    s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = doc.styles[base]
    return s


def _configure_styles(doc, layout: LayoutSpecification):
    body = layout.body_style
    normal = doc.styles["Normal"]
    normal.font.name = body.font_family
    normal.font.size = Pt(body.font_size_pt)
    normal.paragraph_format.space_before = Pt(body.space_before_pt)
    normal.paragraph_format.space_after = Pt(body.space_after_pt)
    normal.paragraph_format.line_spacing = body.line_spacing

    for level, name in ((1, "Heading 1"), (2, "Heading 2"), (3, "Heading 3")):
        style_spec = layout.heading_styles.get(level, body)
        s = doc.styles[name]
        s.base_style = normal
        s.font.name = style_spec.font_family
        s.font.size = Pt(style_spec.font_size_pt)
        s.font.bold = style_spec.bold
        s.font.italic = style_spec.italic
        s.paragraph_format.left_indent = Cm(0)
        s.paragraph_format.space_before = Pt(style_spec.space_before_pt)
        s.paragraph_format.space_after = Pt(style_spec.space_after_pt)
        s.paragraph_format.line_spacing = style_spec.line_spacing

    doan = _style(doc, "DoanVB")
    doan.paragraph_format.first_line_indent = Cm(1)

    tru = _style(doc, "Style Trừ")
    tru.paragraph_format.left_indent = Cm(1)
    tru.paragraph_format.first_line_indent = Cm(-0.5)

    cong = _style(doc, "Style Cộng")
    cong.paragraph_format.left_indent = Cm(1.5)
    cong.paragraph_format.first_line_indent = Cm(-0.5)

    cap = _style(doc, "Caption")
    cap.font.name = layout.caption_style.font_family
    cap.font.size = Pt(layout.caption_style.font_size_pt)
    cap.font.italic = layout.caption_style.italic
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.line_spacing = 1.3


def _ensure_numbering_part(doc):
    from docx.parts.numbering import NumberingPart
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    for rel in doc.part.rels.values():
        if rel.reltype == RT.NUMBERING:
            return rel.target_part

    content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"
    blob = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="{W_NS}">
  <w:abstractNum w:abstractNumId="0">
    <w:multiLevelType w:val="multilevel"/>
    <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="CHƯƠNG %1."/><w:pStyle w:val="Heading1"/></w:lvl>
    <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/><w:pStyle w:val="Heading2"/></w:lvl>
    <w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2.%3."/><w:pStyle w:val="Heading3"/></w:lvl>
  </w:abstractNum>
  <w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>
</w:numbering>""".encode("utf-8")
    part = NumberingPart(PackURI("/word/numbering.xml"), content_type, parse_xml(blob), doc.part.package)
    doc.part.relate_to(part, RT.NUMBERING)
    return part


def _new_numbering(doc):
    root = _ensure_numbering_part(doc).element
    abs_ids = [int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))]
    aid = max(abs_ids or [0]) + 1
    nid = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(aid))
    ml = OxmlElement("w:multiLevelType")
    ml.set(qn("w:val"), "multilevel")
    abstract.append(ml)

    levels = [
        (0, "CHƯƠNG %1.", "Heading1"),
        (1, "%1.%2.", "Heading2"),
        (2, "%1.%2.%3.", "Heading3"),
    ]
    for ilvl, fmt, style in levels:
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, val in [("start", "1"), ("numFmt", "decimal"), ("lvlText", fmt)]:
            x = OxmlElement("w:" + tag)
            x.set(qn("w:val"), val)
            lvl.append(x)
        ps = OxmlElement("w:pStyle")
        ps.set(qn("w:val"), style)
        lvl.append(ps)
        abstract.append(lvl)
    root.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(nid))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(aid))
    num.append(ref)
    root.append(num)
    return nid


def _number(p, nid, level):
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(nid))
    numpr.extend([ilvl, numid])
    ppr.append(numpr)


def _update_fields(doc):
    el = doc.settings._element
    x = el.find(qn("w:updateFields"))
    if x is None:
        x = OxmlElement("w:updateFields")
        el.append(x)
    x.set(qn("w:val"), "true")


def _page_number(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    _clear(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(p, "PAGE", "1")


def _insert_section_before(doc, paragraph):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    sect = section._sectPr
    sect.getparent().remove(sect)
    paragraph._p.addprevious(sect)
    return section


def _caption(doc, anchor, label, text, placeholder="1.1"): 
    p = _after(anchor, "Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label + " ")
    _font(r, italic=True)
    _field(p, f"SEQ {label} \\* ARABIC \\s 1", placeholder)
    r = p.add_run(". " + text)
    _font(r, italic=True)
    return p


def _footnote_reference(p):
    r = p.add_run()
    _font(r)
    rpr = r._r.get_or_add_rPr()
    rs = OxmlElement("w:rStyle")
    rs.set(qn("w:val"), "FootnoteReference")
    rpr.append(rs)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), "1")
    r._r.append(ref)


def _inject_footnote(path):
    temp = path.with_suffix(".tmp.docx")
    path.replace(temp)
    with zipfile.ZipFile(temp) as zin:
        data = {n: zin.read(n) for n in zin.namelist()}

    data["word/footnotes.xml"] = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="{W_NS}">
<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
<w:footnote w:id="1"><w:p><w:r><w:t>Một quốc gia cổ đại tồn tại từ TK 13 đến cuối TK 18</w:t></w:r></w:p></w:footnote>
</w:footnotes>'''.encode()

    relroot = etree.fromstring(data["word/_rels/document.xml.rels"])
    if not any(x.get("Id") == "rIdFootnotes" for x in relroot):
        etree.SubElement(
            relroot,
            f"{{{R_NS}}}Relationship",
            Id="rIdFootnotes",
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
            Target="footnotes.xml",
        )
    data["word/_rels/document.xml.rels"] = etree.tostring(
        relroot, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    ctroot = etree.fromstring(data["[Content_Types].xml"])
    if not any(x.get("PartName") == "/word/footnotes.xml" for x in ctroot):
        etree.SubElement(
            ctroot,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            PartName="/word/footnotes.xml",
            ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        )
    data["[Content_Types].xml"] = etree.tostring(
        ctroot, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in data.items():
            zout.writestr(name, content)
    temp.unlink()


def _clean_heading(value: str) -> str:
    import re
    return re.sub(r"^CHƯƠNG\s+\d+\.\s*|^\d+(?:\.\d+){1,3}\.\s*", "", value).strip()


def build_word_artifact(
    spec: WordTaskSpecification,
    backend: WordBackend | None = None,
) -> Artifact:
    spec.output_dir.mkdir(parents=True, exist_ok=True)
    backend = backend or NativeWordBackend()
    doc = backend.open(spec.template_path or spec.content_source_path)

    from skiptlu.artifacts.template_regions import fingerprint_document, fingerprint_payload
    from skiptlu.artifacts.template_snapshot import snapshot_document, snapshot_payload, validate_snapshot

    template_source = str(spec.template_path or spec.content_source_path or "<generated>")
    template_fingerprint = fingerprint_document(doc, template_source)
    template_snapshot = snapshot_document(doc, template_source)

    task_spec = spec.task_specification
    template_mode = bool(
        task_spec is not None
        and str(task_spec.metadata.get("artifact_strategy", "")) == "use_template"
    )

    if template_mode:
        # Template mode is deliberately isolated from the legacy renderer:
        # the materialized template is the master document and only explicit
        # content operations are allowed to change it.
        from skiptlu.artifacts.template_editor import TemplatePreservingEditor

        TemplatePreservingEditor().edit(
            doc,
            student=spec.student,
            template_fingerprint=template_fingerprint,
            task_operations=task_spec.content_operations if task_spec else [],
            task_contracts=task_spec.region_contracts if task_spec else [],
        )
    else:
        layout = spec.layout_specification or default_word_layout(
            include_toc=True,
            include_landscape=True,
        )
        _configure_styles(doc, layout)
        _update_fields(doc)

        from skiptlu.artifacts.word_renderer import WordRenderer, WordTableRenderer

        numbering_id = _new_numbering(doc)
        renderer = WordRenderer(layout)
        renderer.render(doc, spec, numbering_id)

        # Apply explicit content operations after the legacy renderer.
        if task_spec is not None and task_spec.content_operations:
            from skiptlu.artifacts.content_controller import WordContentController
            WordContentController().apply(
                doc,
                task_spec.content_operations,
                task_spec.region_contracts,
            )

        table_renderer = WordTableRenderer()
        landscape_heading = table_renderer.find_landscape_heading(
            doc,
            next(
                (p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.strip()),
                doc.paragraphs[0],
            ),
        )
        _ = landscape_heading

    output = spec.output_path
    backend.save(doc, output)

    after_snapshot = snapshot_document(doc, template_source)
    structural_validation = validate_snapshot(
        template_snapshot,
        after_snapshot,
        allow_new_paragraphs=not template_mode,
        allow_new_images=not template_mode,
        allow_section_changes=not template_mode,
    )
    if template_mode and structural_validation:
        raise ValueError(
            "Template-preserving validation failed: "
            + "; ".join(structural_validation)
        )
    content = spec.content_specification
    if content and content.include_footnotes and not template_mode:
        _inject_footnote(output)

    return Artifact(
        path=output,
        kind=ArtifactKind.DOCX,
        metadata={
            "activity_id": spec.activity_id,
            "template": str(spec.template_path) if spec.template_path else None,
            "content_source": str(spec.content_source_path) if spec.content_source_path else None,
            "template_fingerprint": fingerprint_payload(template_fingerprint),
            "template_snapshot": snapshot_payload(template_snapshot),
            "post_render_snapshot": snapshot_payload(after_snapshot),
            "template_structural_validation": structural_validation,
        },
    )


def validate_word_artifact(path: Path, spec: WordTaskSpecification) -> list[str]:
    return WordValidator().validate(path, spec)
