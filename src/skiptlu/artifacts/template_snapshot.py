from __future__ import annotations

"""Deterministic snapshot of template structure.

Text is intentionally excluded from the structural signature because content is
allowed to change. Layout, section geometry, paragraph styles, table shape and
embedded-image count are treated as template-owned invariants.
"""

from dataclasses import dataclass
from hashlib import sha256

from docx import Document


@dataclass(frozen=True)
class TemplateSnapshot:
    source: str
    structural_digest: str
    paragraph_count: int
    table_count: int
    section_count: int
    image_count: int
    paragraph_styles: tuple[str, ...]
    table_shapes: tuple[tuple[int, int], ...]
    section_geometry: tuple[tuple[object, ...], ...]


def _structure(doc: Document) -> list[str]:
    parts: list[str] = []

    for index, paragraph in enumerate(doc.paragraphs):
        parts.append(f"P:{index}:{paragraph.style.name}")

    for index, table in enumerate(doc.tables):
        parts.append(f"T:{index}:{len(table.rows)}x{len(table.columns)}")
        for row in table.rows:
            parts.append(f"R:{index}:{len(row.cells)}")

    for index, section in enumerate(doc.sections):
        parts.append(
            "S:"
            f"{index}:{section.orientation}:{section.page_width}:{section.page_height}:"
            f"{section.top_margin}:{section.bottom_margin}:{section.left_margin}:"
            f"{section.right_margin}:{section.header_distance}:{section.footer_distance}"
        )

    parts.append(f"I:{len(doc.inline_shapes)}")
    return parts


def snapshot_document(doc: Document, source: str = "<memory>") -> TemplateSnapshot:
    parts = _structure(doc)
    digest = sha256("\n".join(parts).encode("utf-8")).hexdigest()

    return TemplateSnapshot(
        source=source,
        structural_digest=digest,
        paragraph_count=len(doc.paragraphs),
        table_count=len(doc.tables),
        section_count=len(doc.sections),
        image_count=len(doc.inline_shapes),
        paragraph_styles=tuple(p.style.name for p in doc.paragraphs),
        table_shapes=tuple((len(t.rows), len(t.columns)) for t in doc.tables),
        section_geometry=tuple(
            (
                section.orientation,
                section.page_width,
                section.page_height,
                section.top_margin,
                section.bottom_margin,
                section.left_margin,
                section.right_margin,
                section.header_distance,
                section.footer_distance,
            )
            for section in doc.sections
        ),
    )


def snapshot_payload(snapshot: TemplateSnapshot) -> dict[str, object]:
    return {
        "source": snapshot.source,
        "structural_digest": snapshot.structural_digest,
        "paragraph_count": snapshot.paragraph_count,
        "table_count": snapshot.table_count,
        "section_count": snapshot.section_count,
        "image_count": snapshot.image_count,
        "paragraph_styles": list(snapshot.paragraph_styles),
        "table_shapes": [list(shape) for shape in snapshot.table_shapes],
        "section_geometry": [list(geometry) for geometry in snapshot.section_geometry],
    }


def validate_snapshot(
    before: TemplateSnapshot,
    after: TemplateSnapshot,
    *,
    allow_new_paragraphs: bool = False,
    allow_new_tables: bool = False,
    allow_new_images: bool = False,
    allow_section_changes: bool = False,
) -> list[str]:
    errors: list[str] = []

    if not allow_new_paragraphs and before.paragraph_count != after.paragraph_count:
        errors.append(
            f"Template paragraph count changed: {before.paragraph_count} -> {after.paragraph_count}"
        )

    if not allow_new_tables and before.table_shapes != after.table_shapes:
        errors.append(
            f"Template table structure changed: {before.table_shapes} -> {after.table_shapes}"
        )

    if not allow_new_images and before.image_count != after.image_count:
        errors.append(
            f"Template image count changed: {before.image_count} -> {after.image_count}"
        )

    if not allow_section_changes and before.section_geometry != after.section_geometry:
        errors.append("Template section geometry changed")

    if before.paragraph_styles != after.paragraph_styles:
        errors.append("Template paragraph styles/order changed")

    return errors
