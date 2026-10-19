from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def font(run, size=13, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn("w:" + key), "Times New Roman")


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def replace_paragraph(p, text):
    clear_paragraph(p)
    r = p.add_run(text)
    font(r)
    return r


def add_field(p, instruction, placeholder="1"):
    r = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, separate, t, end])
    font(r)


def paragraph_after(p, style=None):
    new_p = OxmlElement("w:p")
    p._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    new = Paragraph(new_p, p._parent)
    if style:
        new.style = style
    return new


def update_fields(doc):
    element = doc.settings._element
    node = element.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        element.append(node)
    node.set(qn("w:val"), "true")


def page_number(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE", "1")


def insert_section_before(doc, paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.find(qn("w:sectPr"))
    if existing is not None:
        ppr.remove(existing)
    sect = OxmlElement("w:sectPr")
    ppr.append(sect)
    return doc.sections[-1]


def caption(doc, anchor, label, text, placeholder="1.1"):
    p = paragraph_after(anchor, "Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label + " ")
    font(r, italic=True)
    add_field(p, f"SEQ {label} \\* ARABIC \\s 1", placeholder)
    r = p.add_run(". " + text)
    font(r, italic=True)
    return p


def footnote_reference(p):
    r = p.add_run()
    font(r)
    rpr = r._r.get_or_add_rPr()
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "FootnoteReference")
    rpr.append(style)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), "1")
    r._r.append(ref)


def inject_footnote(path: Path):
    temp = path.with_suffix(".tmp.docx")
    path.replace(temp)
    with zipfile.ZipFile(temp) as zin:
        data = {name: zin.read(name) for name in zin.namelist()}

    data["word/footnotes.xml"] = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="{W_NS}">
<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
<w:footnote w:id="1"><w:p><w:r><w:t>Một quốc gia cổ đại tồn tại từ TK 13 đến cuối TK 18</w:t></w:r></w:p></w:footnote>
</w:footnotes>'''.encode()

    relroot = etree.fromstring(data["word/_rels/document.xml.rels"])
    if not any(x.get("Id") == "rIdFootnotes" for x in relroot):
        etree.SubElement(
            relroot,
            f"{{{R_NS}}}Relationship",
            Id="rIdFootnotes",
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
            Target="footnotes.xml",
        )
    data["word/_rels/document.xml.rels"] = etree.tostring(
        relroot, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    ctroot = etree.fromstring(data["[Content_Types].xml"])
    if not any(x.get("PartName") == "/word/footnotes.xml" for x in ctroot):
        etree.SubElement(
            ctroot,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            PartName="/word/footnotes.xml",
            ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        )
    data["[Content_Types].xml"] = etree.tostring(
        ctroot, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in data.items():
            zout.writestr(name, content)
    temp.unlink()
