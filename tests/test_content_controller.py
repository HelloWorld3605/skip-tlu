import pytest
from docx import Document

from skiptlu.artifacts.content_controller import ContentControlError, WordContentController
from skiptlu.tasks.content_operations import ContentOperation, RegionContract


def test_replace_text_is_explicit_and_deterministic():
    doc = Document()
    doc.add_paragraph("Student placeholder: OLD")
    WordContentController().apply(
        doc,
        [ContentOperation(operation="replace_text", target="OLD", value="NEW")],
    )
    assert doc.paragraphs[0].text == "Student placeholder: NEW"


def test_required_missing_target_fails():
    doc = Document()
    doc.add_paragraph("Existing content")
    with pytest.raises(ContentControlError):
        WordContentController().apply(
            doc,
            [ContentOperation(operation="replace_text", target="MISSING", value="X")],
        )


def test_optional_missing_target_is_ignored():
    doc = Document()
    doc.add_paragraph("Existing content")
    WordContentController().apply(
        doc,
        [ContentOperation(
            operation="replace_text",
            target="MISSING",
            value="X",
            required=False,
        )],
    )
    assert doc.paragraphs[0].text == "Existing content"


def test_region_contract_allows_declared_operation():
    doc = Document()
    doc.add_paragraph("Heading", style="Heading 1")
    WordContentController().apply(
        doc,
        [ContentOperation(
            operation="replace_paragraph",
            region_id="body.heading1.0",
            value="Edited",
        )],
        [RegionContract(
            region_id="body.heading1.0",
            allowed_operations={"replace_paragraph"},
        )],
    )
    assert doc.paragraphs[0].text == "Edited"


def test_region_contract_rejects_undeclared_operation():
    doc = Document()
    doc.add_paragraph("Heading", style="Heading 1")
    with pytest.raises(ContentControlError, match="not allowed"):
        WordContentController().apply(
            doc,
            [ContentOperation(
                operation="delete_paragraph",
                region_id="body.heading1.0",
            )],
            [RegionContract(
                region_id="body.heading1.0",
                allowed_operations={"replace_paragraph"},
            )],
        )


def test_region_contract_rejects_unknown_region():
    doc = Document()
    doc.add_paragraph("Heading", style="Heading 1")
    with pytest.raises(ContentControlError, match="No region contract"):
        WordContentController().apply(
            doc,
            [ContentOperation(
                operation="replace_paragraph",
                region_id="body.heading1.99",
                value="Edited",
            )],
            [RegionContract(
                region_id="body.heading1.0",
                allowed_operations={"replace_paragraph"},
            )],
        )


def test_region_contract_forbids_unscoped_mutation():
    doc = Document()
    doc.add_paragraph("OLD")
    with pytest.raises(ContentControlError, match="Unscoped"):
        WordContentController().apply(
            doc,
            [ContentOperation(operation="replace_text", target="OLD", value="NEW")],
            [RegionContract(
                region_id="body.heading1.0",
                allowed_operations={"replace_paragraph"},
            )],
        )


def test_insert_and_delete_are_supported():
    doc = Document()
    doc.add_paragraph("Anchor")
    WordContentController().apply(
        doc,
        [ContentOperation(operation="insert_after", target="Anchor", value="Inserted")],
    )
    assert [p.text for p in doc.paragraphs] == ["Anchor", "Inserted"]

    WordContentController().apply(
        doc,
        [ContentOperation(operation="delete_paragraph", target="Inserted")],
    )
    assert [p.text for p in doc.paragraphs] == ["Anchor"]
