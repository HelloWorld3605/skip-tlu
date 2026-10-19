from pathlib import Path

from docx import Document

from skiptlu.artifacts.word import build_word_artifact, validate_word_artifact
from skiptlu.artifacts.word_backend import NativeWordBackend
from skiptlu.models.word_task import StudentInfo, WordTaskSpecification
from skiptlu.tasks.content import ContentFigure, ContentSection, ContentSpecification, ContentTable
from skiptlu.tasks.content_operations import ContentOperation, RegionContract
from skiptlu.tasks.layout import default_word_layout
from skiptlu.tasks.specification import ArtifactType, PracticalTaskSpecification


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "artifacts" / "reference" / "BaiLam_LuyenTap_BaiSo1.docx"
IMAGE = ROOT / "artifacts" / "reference" / "Picture_De1.jpg"


class RecordingWordBackend(NativeWordBackend):
    name = "recording-native"

    def __init__(self):
        self.opened = []
        self.saved = []

    def open(self, template_path: Path) -> Document:
        self.opened.append(template_path)
        return super().open(template_path)

    def save(self, document: Document, output_path: Path) -> None:
        self.saved.append(output_path)
        super().save(document, output_path)


def make_spec(tmp_path):
    content = ContentSpecification(
        title="Chiang Mai",
        sections=[
            ContentSection(level=1, title="GIỚI THIỆU VỀ CHIANG MAI"),
            ContentSection(level=1, title="CẨM NANG DU LỊCH CHIANG MAI"),
            ContentSection(level=2, title="Thời điểm du lịch lý tưởng"),
            ContentSection(level=2, title="Phương tiện"),
            ContentSection(level=2, title="Địa điểm thăm quan"),
            ContentSection(level=3, title="Phương tiện đến Chiang Mai"),
            ContentSection(level=3, title="Các phương tiện di chuyển tại Chiang Mai"),
        ],
        figures=[ContentFigure(label="Ảnh", caption="Ảnh 1.1. Chiang Mai mùa hoa nở")],
        tables=[ContentTable(label="Bảng", caption="Bảng 2.1. Những điểm thăm quan nổi tiếng")],
        include_toc=True,
        include_landscape_section=True,
    )
    return WordTaskSpecification(
        activity_id=51068,
        title="Chiang Mai",
        template_path=TEMPLATE,
        image_path=IMAGE,
        output_dir=tmp_path,
        student=StudentInfo(
            name="TEST STUDENT",
            student_id="TEST000",
            major="TÀI CHÍNH - NGÂN HÀNG",
        ),
        task_specification=PracticalTaskSpecification(
            activity_id=51068,
            title="Chiang Mai",
            artifact_type=ArtifactType.DOCX,
            content=content,
            layout=default_word_layout(include_toc=True, include_landscape=True),
        ),
        content_specification=content,
        layout_specification=default_word_layout(include_toc=True, include_landscape=True),
    )


def test_week4_word_artifact_from_real_template(tmp_path):
    spec = make_spec(tmp_path)

    artifact = build_word_artifact(spec)
    assert artifact.path.exists()

    errors = validate_word_artifact(artifact.path, spec)
    assert errors == [], errors

    doc = Document(str(artifact.path))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "SINH VIÊN THỰC HIỆN: TEST STUDENT" in body
    assert "MÃ SINH VIÊN: TEST000" in body
    assert len(doc.sections) >= 2
    assert len(doc.tables) == 1
    assert len(doc.inline_shapes) >= 2


def test_content_operations_are_applied_after_template_render(tmp_path):
    spec = make_spec(tmp_path)
    spec.task_specification.content_operations = [
        ContentOperation(
            operation="replace_text",
            target="TEST STUDENT",
            value="EDITED STUDENT",
        )
    ]

    artifact = build_word_artifact(spec)

    doc = Document(str(artifact.path))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "EDITED STUDENT" in body
    assert "TEST STUDENT" not in body
    assert validate_word_artifact(artifact.path, spec) == []


def test_region_operation_targets_template_structure(tmp_path):
    spec = make_spec(tmp_path)
    spec.task_specification.content_operations = [
        ContentOperation(
            operation="replace_paragraph",
            region_id="body.heading1.0",
            value="EDITED HEADING",
        )
    ]
    spec.task_specification.region_contracts = [
        RegionContract(
            region_id="body.heading1.0",
            allowed_operations={"replace_paragraph"},
        )
    ]

    artifact = build_word_artifact(spec)
    doc = Document(str(artifact.path))
    assert any(p.text == "EDITED HEADING" for p in doc.paragraphs)
    assert "template_fingerprint" in artifact.metadata
    assert artifact.metadata["template_fingerprint"]["regions"]


def test_word_backend_is_injected_for_document_io(tmp_path):
    spec = make_spec(tmp_path)
    backend = RecordingWordBackend()

    artifact = build_word_artifact(spec, backend=backend)

    assert backend.opened == [TEMPLATE]
    assert backend.saved == [artifact.path]
    assert artifact.path.exists()
    assert validate_word_artifact(artifact.path, spec) == []
