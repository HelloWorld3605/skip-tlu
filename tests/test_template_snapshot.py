from docx import Document

from skiptlu.artifacts.template_snapshot import (
    snapshot_document,
    validate_snapshot,
)


def test_snapshot_ignores_allowed_text_changes():
    doc = Document()
    doc.add_paragraph("OLD", style="Heading 1")
    before = snapshot_document(doc)

    doc.paragraphs[0].text = "NEW"
    after = snapshot_document(doc)

    assert before.structural_digest == after.structural_digest
    assert validate_snapshot(before, after) == []


def test_snapshot_detects_paragraph_structure_change():
    doc = Document()
    doc.add_paragraph("One")
    before = snapshot_document(doc)

    doc.add_paragraph("Two")
    after = snapshot_document(doc)

    errors = validate_snapshot(before, after)
    assert any("paragraph count" in error for error in errors)


def test_snapshot_detects_section_geometry_change():
    doc = Document()
    before = snapshot_document(doc)

    doc.sections[0].left_margin = 123456
    after = snapshot_document(doc)

    errors = validate_snapshot(before, after)
    assert "Template section geometry changed" in errors


def test_snapshot_detects_table_shape_change():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    before = snapshot_document(doc)

    table.add_row()
    after = snapshot_document(doc)

    errors = validate_snapshot(before, after)
    assert any("table structure" in error for error in errors)
